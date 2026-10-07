"""Tests for inventory backend (app/inventory.py, app/exporter.py, API)."""

from __future__ import annotations

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.inventory import categorize_product_type, validate_numeric
import pytest


def test_categorize_product_type_keywords():
    assert categorize_product_type("Mouse sem fio") == "Eletrônicos"
    assert categorize_product_type("ARROZ 5kg") == "Alimentos"
    assert categorize_product_type("Detergente neutro") == "Limpeza"
    assert categorize_product_type("Caneta azul") == "Escritório"
    assert categorize_product_type("Martelo 500g") == "Ferramentas"


def test_categorize_product_type_fallback():
    assert categorize_product_type("Coisa misteriosa xyz") == "Outros"
    assert categorize_product_type("") == "Outros"


def test_validate_numeric_ok():
    assert validate_numeric("10", "quantity") == 10.0
    assert validate_numeric(2.5, "unit_price") == 2.5


def test_validate_numeric_rejects_non_numeric():
    for bad in ("abc", None, True, ""):
        with pytest.raises(ValueError):
            validate_numeric(bad, "quantity")


def test_post_item_autocategorizes(client: TestClient):
    resp = client.post(
        "/estoque/itens",
        json={"product_name": "Mouse sem fio", "quantity": 5, "unit_price": 49.9},
    )
    assert resp.status_code == 201
    body = resp.json()
    assert body["product_type"] == "Eletrônicos"
    assert body["quantity"] == 5.0


def test_post_item_rejects_non_numeric(client: TestClient):
    resp = client.post(
        "/estoque/itens",
        json={"product_name": "Mouse", "quantity": "abc", "unit_price": 1.0},
    )
    assert resp.status_code == 422


def test_post_item_rejects_negative(client: TestClient):
    resp = client.post(
        "/estoque/itens",
        json={"product_name": "Mouse", "quantity": -1, "unit_price": 1.0},
    )
    assert resp.status_code == 422


def test_items_grouped_by_type(client: TestClient):
    client.post("/estoque/itens", json={"product_name": "Mouse", "quantity": 2, "unit_price": 10.0})
    client.post("/estoque/itens", json={"product_name": "Arroz 5kg", "quantity": 3, "unit_price": 20.0})
    client.post("/estoque/itens", json={"product_name": "Teclado USB", "quantity": 1, "unit_price": 80.0})
    resp = client.get("/estoque/grupos")
    assert resp.status_code == 200
    groups = {g["product_type"]: g for g in resp.json()}
    assert groups["Eletrônicos"]["items"] == 2
    assert groups["Alimentos"]["items"] == 1


def test_export_excel_one_click(client: TestClient):
    client.post("/estoque/itens", json={"product_name": "Mouse", "quantity": 2, "unit_price": 10.0})
    resp = client.get("/estoque.xlsx")
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers["content-type"]
    assert "estoque.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(filename=__import__("io").BytesIO(resp.content))
    assert "Estoque" in wb.sheetnames
    assert "Por tipo" in wb.sheetnames
    ws = wb["Estoque"]
    assert ws.max_row == 2  # header + 1 item
    assert ws.cell(row=1, column=2).value == "Produto"


def test_estoque_page_served(client: TestClient):
    resp = client.get("/estoque")
    assert resp.status_code == 200
    assert "Registrar item" in resp.text
