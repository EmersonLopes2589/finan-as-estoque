"""Excel export for inventory data.

Builds a ``.xlsx`` workbook (via openpyxl) with two sheets:
``Estoque`` (all items) and ``Por tipo`` (aggregated totals).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font

ITEM_HEADERS: tuple[str, ...] = (
    "ID",
    "Produto",
    "Tipo",
    "Quantidade",
    "Preço unitário",
    "Registrado em",
)

GROUP_HEADERS: tuple[str, ...] = ("Tipo", "Itens", "Quantidade total")


def _autosize(ws) -> None:
    """Adjust column widths to fit content."""
    for column in ws.columns:
        width = max(
            (len(str(cell.value)) if cell.value is not None else 0 for cell in column),
            default=0,
        )
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 40)


def build_inventory_workbook(
    items: Sequence[Mapping[str, Any]],
    groups: Sequence[Mapping[str, Any]],
) -> bytes:
    """Return ``.xlsx`` bytes with all *items* plus per-type totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Estoque"
    ws.append(list(ITEM_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for item in items:
        ws.append(
            [
                item.get("id"),
                item.get("product_name"),
                item.get("product_type"),
                item.get("quantity"),
                item.get("unit_price"),
                item.get("created_at"),
            ]
        )
    ws.freeze_panes = "A2"
    _autosize(ws)

    ws2 = wb.create_sheet("Por tipo")
    ws2.append(list(GROUP_HEADERS))
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for group in groups:
        ws2.append(
            [group.get("product_type"), group.get("items"), group.get("total_quantity")]
        )
    ws2.freeze_panes = "A2"
    _autosize(ws2)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
